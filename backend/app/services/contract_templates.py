"""Original-template catalog, retrieval and reviewed structural DOCX filling."""

import re
import uuid
from dataclasses import dataclass
from functools import lru_cache
from hashlib import sha256
from pathlib import Path

from app.documents.field_metadata import field_metadata
from app.documents.slot_constraints import validate_slot_values
from app.documents.template_slots import SLOT_ENGINE_VERSION, compile_template_slots
from app.documents.word_compatibility import iter_visible_document_paragraphs
from app.documents.word_documents import iter_document_paragraphs
from app.documents.word_integrity import visual_fingerprint
from app.documents.word_text import paragraph_text, replace_text_ranges
from app.domain.errors import InvalidContractInput
from app.services.template_retrieval import (
    RetrievalPolicy,
    explicit_source_titles,
    normalize_title,
    rank_metadata,
    title_aliases,
)

GROUPS = {"Ministerial": "部委", "Regional": "地方"}


@dataclass(frozen=True)
class ContractTemplate:
    id: str
    name: str
    group: str
    path: Path

    def public(self) -> dict:
        return {"id": self.id, "name": self.name, "group": GROUPS[self.group]}


class ContractTemplateCatalog:
    def __init__(
        self,
        resources: Path,
        generated: Path,
        embedding_path: Path | None = None,
        retrieval_policy: RetrievalPolicy | None = None,
    ):
        self.resources = resources
        self.generated = generated
        self.embedding_path = embedding_path
        self.retrieval_policy = retrieval_policy or RetrievalPolicy()
        self._templates: dict[str, ContractTemplate] = {}

    def refresh(self) -> None:
        """Only files discovered beneath the configured template root can be selected."""
        templates = {}
        # Only untouched originals are a production catalog source.
        root = (self.resources / "templates").resolve()
        for group in GROUPS:
            folder = root / group
            if not folder.is_dir():
                continue
            for path in folder.glob("*.docx"):
                if not path.resolve().is_relative_to(root):
                    continue
                identifier = sha256(f"{group}/{path.stem}".encode()).hexdigest()[:20]
                templates[identifier] = ContractTemplate(
                    identifier,
                    path.stem,
                    group,
                    path,
                )
        self._templates = templates

    def list_templates(self, group: str | None = None, query: str = "", limit: int = 30) -> list[dict]:
        if not self._templates:
            self.refresh()
        items = self._templates.values()
        if group:
            normalized = next((key for key, label in GROUPS.items() if group in (key, label)), None)
            if normalized is None:
                raise InvalidContractInput("未知合同分类")
            items = (item for item in items if item.group == normalized)
        if query:
            items = (item for item in items if query.casefold() in item.name.casefold())
        return [
            item.public()
            for item in sorted(items, key=lambda item: (item.group, item.name))[: max(1, min(limit, 100))]
        ]

    def browse_templates(
        self, group: str | None = None, query: str = "", offset: int = 0, limit: int = 24
    ) -> dict:
        """Paginate the full catalog so manual choice never depends on Agent ranking."""
        if not self._templates:
            self.refresh()
        items = list(self._templates.values())
        if group:
            normalized = next((key for key, label in GROUPS.items() if group in (key, label)), None)
            if normalized is None:
                raise InvalidContractInput("未知合同分类")
            items = [item for item in items if item.group == normalized]
        if query.strip():
            terms = query.strip().casefold().split()
            items = [
                item
                for item in items
                if all(term in (item.name + self._keywords(item)).casefold() for term in terms)
            ]
        items.sort(key=lambda item: (item.group, item.name))
        return {
            "items": [item.public() for item in items[offset : offset + limit]],
            "total": len(items),
            "offset": offset,
            "limit": limit,
        }

    def get_template(self, template_id: str) -> ContractTemplate:
        if not self._templates:
            self.refresh()
        template = self._templates.get(template_id)
        if template is None:
            raise InvalidContractInput("合同模板不存在")
        return template

    def describe_template(self, template_id: str) -> dict:
        template = self.get_template(template_id)
        document, compilation = self._compiled_document(template)
        paragraph_objects = list(iter_visible_document_paragraphs(document))
        paragraphs = tuple(paragraph_text(paragraph) for paragraph in paragraph_objects)
        contexts: dict[str, list[str]] = {}
        metadata = {}
        for paragraph_object in paragraph_objects:
            paragraph = paragraph_text(paragraph_object)
            for match in re.finditer(r"\{([^{}]+)\}", paragraph):
                field_name = match.group(1)
                # Mask neighbours before clipping, otherwise a sliced opaque
                # locator can leak into the prompt and obscure the active slot.
                marked = re.sub(
                    r"\{([^{}]+)\}",
                    lambda item: item.group(0) if item.group(1) == field_name else "[其他填写处]",
                    paragraph,
                )
                active = marked.index("{" + field_name + "}")
                context = marked[max(0, active - 150) : active + len(field_name) + 152]
                if paragraph_object._p.xpath("ancestor::w:txbxContent"):
                    context = "文本框：" + context
                cell = paragraph_object._p.getparent()
                if cell.tag.endswith("}tc"):
                    table = cell.getparent().getparent()
                    rows = table.xpath("./w:tr")
                    header = (
                        " | ".join("".join(item.xpath(".//w:t/text()")) for item in rows[0].xpath("./w:tc"))
                        if rows
                        else ""
                    )
                    row_context = " | ".join(
                        "".join(item.xpath(".//w:t/text()")) for item in cell.getparent().xpath("./w:tc")
                    )
                    column = len(list(cell.itersiblings(preceding=True, tag=cell.tag))) + 1
                    context += f"\n表格第{column}列；首行：{header[:1500]}\n当前行：{row_context[:1500]}"
                    row_number = next(
                        (index + 1 for index, row in enumerate(rows) if row is cell.getparent()), 0
                    )
                    headings = (
                        ["".join(item.xpath(".//w:t/text()")).strip() for item in rows[0].xpath("./w:tc")]
                        if rows
                        else []
                    )
                    heading = re.sub(
                        r"\{[^{}]+\}", "", headings[column - 1] if column <= len(headings) else ""
                    )
                    metadata[field_name] = {
                        "display_name": f"第{row_number}行 · {heading or '第' + str(column) + '列'}",
                        "section": "表格明细",
                        "kind": "table",
                    }
                if context not in contexts.setdefault(field_name, []):
                    contexts[field_name].append(context)
        # The compiled Word structure is authoritative; legacy JSON is not read.
        fields = [
            {
                "name": name,
                "context": "\n".join(examples[:3]),
                **field_metadata(name, examples[0]),
                **compilation.field_specs.get(name, {}),
                **metadata.get(name, {}),
            }
            for name, examples in contexts.items()
        ]
        preview_text = "\n".join(paragraphs)
        warnings = [
            "填写位置由结构扫描识别，不等同于模板业务完整性认证；请下载原件检查未识别空白、选项及附件。"
        ]
        if compilation.generated_slots:
            warnings.append(
                f"已从原始文档识别 {compilation.generated_slots} 处候选空白；排版空格或非填写区域可注明原因后留空。"
            )
        if compilation.unsupported:
            warnings.append("存在不支持的文档结构，禁止自动填写：" + "、".join(compilation.unsupported))
        warnings.extend(compilation.warnings)
        return {
            **template.public(),
            "placeholders": list(contexts),
            "fields": fields,
            "preview_text": preview_text[:60000],
            "preview_truncated": len(preview_text) > 60000,
            "warnings": warnings,
            "fillable": bool(fields) and not compilation.unsupported,
            "unsupported_structures": list(compilation.unsupported),
            "generated_slots": compilation.generated_slots,
            "requires_visual_review": compilation.requires_visual_review,
            "schema_fingerprint": sha256(
                (SLOT_ENGINE_VERSION + sha256(template.path.read_bytes()).hexdigest()).encode()
            ).hexdigest(),
        }

    def preview_contract(
        self, template_id: str, values: dict[str, str], omitted_fields: list[str] | None = None
    ) -> dict:
        template = self.get_template(template_id)
        document, compilation = self._compiled_document(template)
        paragraph_objects = tuple(iter_visible_document_paragraphs(document))
        paragraphs = tuple(paragraph_text(paragraph) for paragraph in paragraph_objects)
        replacements = {
            name: compilation.original_blanks[name]
            for name in omitted_fields or []
            if name in compilation.original_blanks
        }
        replacements.update(values)
        validate_slot_values(self.describe_template(template_id)["fields"], values)
        rendered = [
            re.sub(r"\{([^{}]+)\}", lambda match: replacements.get(match.group(1), match.group(0)), paragraph)
            for paragraph in paragraphs
        ]
        text = "\n".join(rendered)
        blocks, tables = [], {}
        for paragraph, rendered_text in zip(paragraph_objects, rendered):
            if not rendered_text.strip():
                continue
            cell = paragraph._p.getparent()
            if cell.tag.endswith("}tc"):
                row, table = cell.getparent(), cell.getparent().getparent()
                if table not in tables:
                    block = {"type": "table", "rows": []}
                    tables[table] = (block, {})
                    blocks.append(block)
                block, row_map = tables[table]
                if row not in row_map:
                    # Preserve blank columns; otherwise a value can appear under the wrong header.
                    row_map[row] = {item: "" for item in row.xpath("./w:tc")}
                    block["rows"].append(row_map[row])
                row_map[row][cell] = (row_map[row].get(cell, "") + "\n" + rendered_text).strip()
            else:
                blocks.append({"type": "paragraph", "text": rendered_text})
        for block in blocks:
            if block["type"] == "table":
                block["rows"] = [list(row.values()) for row in block["rows"]]
        # Very large previews use the bounded text fallback instead of an unbounded block payload.
        return {
            "text": text[:120000],
            "blocks": blocks if len(text) <= 120000 else [],
            "truncated": len(text) > 120000,
        }

    @staticmethod
    def _compiled_document(template: ContractTemplate):
        from docx import Document

        document = Document(template.path)
        before = visual_fingerprint(document)
        compilation = compile_template_slots(document)
        if visual_fingerprint(document) != before:
            raise InvalidContractInput("结构编译改变了图形或媒体内容，已阻止填写")
        return document, compilation

    def search_templates(
        self,
        query: str,
        group: str | None = None,
        limit: int = 5,
        *,
        explicit_names: list[str] | None = None,
        source_text: str = "",
    ) -> list[dict]:
        """Retrieve up to N usable candidates, not claim a globally optimal N."""
        if not query.strip():
            raise InvalidContractInput("请描述合同需求")
        if not self._templates:
            self.refresh()
        candidates = list(self._templates.values())
        if group:
            normalized = next((key for key, label in GROUPS.items() if group in (key, label)), None)
            if normalized is None:
                raise InvalidContractInput("未知合同分类")
            candidates = [item for item in candidates if item.group == normalized]
        vector_candidates = self._vector_candidates(query, group)
        records = [
            {"id": item.id, "name": item.name, "keywords": self._keywords(item)} for item in candidates
        ]
        vector_ranks = {
            item.id: vector_candidates[(item.group, item.name)]
            for item in candidates
            if (item.group, item.name) in vector_candidates
        }
        ranked = rank_metadata(query, records, vector_ranks, self.retrieval_policy)
        requested = {normalize_title(name) for name in explicit_names or []}
        requested.update(explicit_source_titles(source_text, [item.name for item in candidates]))
        full_titles = {normalize_title(item.name) for item in candidates} & requested
        requested = {
            name for name in requested if not any(name != full and name in full for full in full_titles)
        }
        named_ids = {item.id for item in candidates if title_aliases(item.name) & requested}
        ranked_by_id = {item["id"]: item for item in ranked}
        named = [
            ranked_by_id.get(identifier, {"id": identifier, "score": 1.0, "sources": []})
            for identifier in sorted(named_ids, key=lambda key: (self._templates[key].name, key))
        ]
        ranked = named + [item for item in ranked if item["id"] not in named_ids]
        target = max(1, min(limit, 10))
        recommendations = []
        labels = {
            "title_bm25": "名称 BM25 词频与稀有词匹配",
            "keyword_bm25": "业务关键词 BM25 匹配",
            "vector": "可选语义索引召回",
            "exact_title": "名称完整片段匹配",
        }
        for item in ranked[: max(target, self.retrieval_policy.candidate_pool)]:
            template = self.get_template(item["id"])
            try:
                details = self.describe_template(template.id)
            except (InvalidContractInput, OSError, ValueError):
                continue
            explicitly_requested = template.id in named_ids
            if not details["fillable"] and not explicitly_requested:
                continue
            reasons = [labels[source] for source in item["sources"]]
            if explicitly_requested:
                reasons.insert(0, "会话明确点名的合同，优先展示；仍请核对版本及最终约定")
            if details["fillable"]:
                reasons.append("通过当前结构填写检查；不等于业务适用性认证")
            else:
                reasons.append(
                    "已找到原件，但暂不能自动填写："
                    + "、".join(details["unsupported_structures"] or ["未识别到字段"])
                )
            if details["requires_visual_review"]:
                reasons.append("含保留图形，导出前需额外人工确认")
            recommendations.append(
                {
                    **template.public(),
                    "match_score": item["score"],
                    "match_reasons": reasons,
                    "retrieval_mode": "bm25_vector_rrf" if vector_ranks else "bm25_rrf",
                    "retrieval_sources": item["sources"],
                    "fillable": details["fillable"],
                    "explicitly_requested": explicitly_requested,
                    "requires_visual_review": details["requires_visual_review"],
                }
            )
            if len(recommendations) == target:
                break
        return recommendations

    def _keywords(self, template: ContractTemplate) -> str:
        path = self.resources / "templates" / template.group / "关键词" / f"{template.name}.txt"
        return self._read_keywords(str(path))

    @staticmethod
    @lru_cache(maxsize=1024)
    def _read_keywords(path: str) -> str:
        try:
            return Path(path).read_text(encoding="utf-8")[:1000]
        except OSError:
            return ""

    def _vector_candidates(self, query: str, group: str | None) -> dict[tuple[str, str], float]:
        if not self.embedding_path:
            return {}
        try:
            import faiss

            model = self._embedding_model(str(self.embedding_path))
            vector = model.encode([query]).astype("float32")
            result = {}
            for key in GROUPS:
                if group and group not in (key, GROUPS[key]):
                    continue
                folder = self.resources / "indexes" / key
                names = [
                    Path(line.strip()).stem
                    for line in (folder / "filenames.txt").read_text(encoding="utf-8").splitlines()
                    if line.strip()
                ]
                index = faiss.read_index(str(folder / "knowledge_base.index"))
                if index.ntotal != len(names):
                    continue
                distances, indices = index.search(vector, min(20, len(names)))
                for rank, position in enumerate(indices[0]):
                    if 0 <= position < len(names):
                        result[(key, names[position])] = 1.0 - rank / max(1, len(indices[0]))
            return result
        except (ImportError, OSError, ValueError, RuntimeError):
            return {}

    @staticmethod
    @lru_cache(maxsize=2)
    def _embedding_model(path: str):
        from sentence_transformers import SentenceTransformer

        return SentenceTransformer(path, local_files_only=True)

    def generate_contract(
        self, template_id: str, values: dict[str, str], omitted_fields: list[str] | None = None
    ) -> dict:
        """Write reviewed locations, checking the resulting text before saving."""
        template = self.get_template(template_id)
        details = self.describe_template(template_id)
        if not details["fillable"]:
            raise InvalidContractInput("此模板存在不支持的结构或没有可填写位置，请维护模板后使用")
        allowed = set(details["placeholders"])
        clean = {
            key: str(value).strip()
            for key, value in values.items()
            if key in allowed and value is not None and str(value).strip()
        }
        if not clean:
            raise InvalidContractInput("未提取到可填写的合同字段，请补充当事人、金额等具体信息")
        validate_slot_values(details["fields"], clean)
        document, compilation = self._compiled_document(template)
        visual_before = visual_fingerprint(document)
        replacements = {
            name: compilation.original_blanks[name]
            for name in omitted_fields or []
            if name in compilation.original_blanks
        }
        replacements.update(clean)
        filled = set()
        for paragraph in iter_document_paragraphs(document):
            expected = re.sub(
                r"\{([^{}]+)\}",
                lambda match: replacements.get(match.group(1), match.group(0)),
                paragraph_text(paragraph),
            )
            filled.update(self._fill_paragraph(paragraph, replacements))
            if paragraph_text(paragraph) != expected:
                raise InvalidContractInput("模板写入结果与已审核内容不一致，已停止生成，请维护模板结构")
        if visual_fingerprint(document) != visual_before:
            raise InvalidContractInput("填写改变了图形、文本框框体或媒体内容，已阻止生成")
        if set(replacements) - filled:
            raise InvalidContractInput("部分已审核的填写或留空位置未能写入，已停止生成")
        filled &= set(clean)
        if not filled:
            raise InvalidContractInput("模板中未找到对应填写位置，无法生成合同")
        if set(clean) - filled:
            raise InvalidContractInput(
                "部分已确认字段无法写入模板，请更换模板或检查文档结构："
                + "、".join(sorted(set(clean) - filled)[:5])
            )
        self.generated.mkdir(parents=True, exist_ok=True)
        contract_id = uuid.uuid4().hex
        filename = f"{contract_id}.docx"
        document.save(self.generated / filename)
        return {
            "id": contract_id,
            "template": template.public(),
            "filled_fields": sorted(filled),
            "missing_fields": sorted(allowed - filled),
            "filename": filename,
        }

    @staticmethod
    def _fill_paragraph(paragraph, values: dict[str, str]) -> set[str]:
        """Write only text leaves, including split fields, never flatten a run."""
        original = paragraph_text(paragraph)
        matches = [match for match in re.finditer(r"\{([^{}]+)\}", original) if match.group(1) in values]
        replace_text_ranges(
            paragraph, [(match.start(), match.end(), values[match.group(1)]) for match in matches]
        )
        return {match.group(1) for match in matches}
