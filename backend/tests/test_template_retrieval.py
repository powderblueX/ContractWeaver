"""Retrieval must be deterministic, bounded and honest about candidate ranking."""

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document
from docx.oxml import OxmlElement

from app.services.contract_templates import ContractTemplateCatalog
from app.services.template_retrieval import (
    RetrievalPolicy,
    bm25_scores,
    explicit_source_titles,
    normalize_title,
    rank_metadata,
    tokenize,
)


class TemplateRetrievalTests(TestCase):
    """测试夹具或替身：TemplateRetrievalTests，用于隔离验证本模块行为。"""

    def test_literal_requests_survive_model_omission_but_negation_and_questions_do_not(self):
        """Verify literal requests survive model omission but negation and questions do not."""
        names = ["设备采购合同", "房屋租赁合同"]
        self.assertEqual(explicit_source_titles("需要《设备采购合同》。", names), {"设备采购合同"})
        self.assertEqual(explicit_source_titles("不需要设备采购合同，房屋租赁合同需要吗？", names), set())
        self.assertEqual(normalize_title("《设备采购合同》 （2025）"), "设备采购合同(2025)")

    def test_named_contract_is_prioritized_even_if_unrelated_query_or_blocked(self):
        """Verify named contract is prioritized even if unrelated query or blocked."""
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            folder = root / "templates" / "Ministerial"
            folder.mkdir(parents=True)
            for name in ["设备采购合同", "房屋租赁合同"]:
                document = Document()
                document.add_paragraph("甲方____")
                if name == "房屋租赁合同":
                    document.element.body.append(OxmlElement("w:sdt"))
                document.save(folder / f"{name}.docx")
            catalog = ContractTemplateCatalog(root, root / "generated")
            recommendations = catalog.search_templates("设备采购", source_text="需要《房屋租赁合同》。")
            self.assertEqual(recommendations[0]["name"], "房屋租赁合同")
            self.assertTrue(recommendations[0]["explicitly_requested"])
            self.assertFalse(recommendations[0]["fillable"])
            self.assertIn("暂不能自动填写", "".join(recommendations[0]["match_reasons"]))

    def test_specific_version_does_not_pin_other_versions(self):
        """Verify specific version does not pin other versions."""
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            folder = root / "templates" / "Ministerial"
            folder.mkdir(parents=True)
            for year in (2020, 2025):
                document = Document()
                document.add_paragraph("甲方____")
                document.save(folder / f"设备采购合同（{year}）.docx")
            catalog = ContractTemplateCatalog(root, root / "generated")
            recommendations = catalog.search_templates("无关内容", source_text="采用设备采购合同（2025）。")
            self.assertEqual([item["name"] for item in recommendations], ["设备采购合同（2025）"])

    def test_tokenization_supports_chinese_and_latin_identifiers(self):
        """Verify tokenization supports chinese and latin identifiers."""
        self.assertEqual(tokenize("设备采购 GF-2025"), ["设备", "备采", "采购", "gf", "2025"])

    def test_bm25_uses_rare_terms_instead_of_only_overlap_counts(self):
        """Verify bm25 uses rare terms instead of only overlap counts."""
        scores = bm25_scores("设备 质谱", ["设备采购合同", "设备租赁合同", "质谱设备采购"], RetrievalPolicy())
        self.assertGreater(scores[2], scores[0])
        self.assertGreater(scores[2], scores[1])

    def test_rrf_uses_both_lexical_and_vector_ranks_without_unknown_ids(self):
        """Verify rrf uses both lexical and vector ranks without unknown ids."""
        records = [
            {"id": "one", "name": "设备采购合同", "keywords": "采购"},
            {"id": "two", "name": "机械买卖合同", "keywords": "设备采购"},
        ]
        ranked = rank_metadata("设备采购", records, {"two": 1.0, "unknown": 99}, RetrievalPolicy())
        self.assertEqual({item["id"] for item in ranked}, {"one", "two"})
        second = next(item for item in ranked if item["id"] == "two")
        self.assertIn("vector", second["sources"])
        self.assertIn("keyword_bm25", second["sources"])
        self.assertEqual(
            ranked, rank_metadata("设备采购", records, {"two": 1.0, "unknown": 99}, RetrievalPolicy())
        )

    def test_catalog_recommends_only_usable_matches_and_does_not_pad_to_five(self):
        """Verify catalog recommends only usable matches and does not pad to five."""
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            folder = root / "templates" / "Ministerial"
            folder.mkdir(parents=True)
            good = Document()
            good.add_paragraph("甲方____")
            good.save(folder / "设备采购合同.docx")
            blocked = Document()
            blocked.add_paragraph("甲方____")
            blocked.element.body.append(OxmlElement("w:sdt"))
            blocked.save(folder / "设备采购合同特别版.docx")
            catalog = ContractTemplateCatalog(root, root / "generated")
            recommendations = catalog.search_templates("设备采购", limit=5)
            self.assertEqual([item["name"] for item in recommendations], ["设备采购合同"])
            self.assertEqual(recommendations[0]["retrieval_mode"], "bm25_rrf")
            self.assertEqual(catalog.search_templates("完全不相关的保险"), [])
