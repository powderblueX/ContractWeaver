"""End-to-end workflow checks using real SQLite/DOCX and deterministic agent roles."""

import asyncio
import json
import os
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from docx.oxml import parse_xml
from fastapi.testclient import TestClient

from app.agents.team import create_agent_team as real_agent_team
from app.api.dependencies import get_database
from app.api.workspace_dependencies import get_template_catalog
from app.config import Settings, get_settings
from app.domain.contract_workflow import (
    CoverageResult,
    ExtractionResult,
    NegotiationAnalysis,
    ReviewResult,
    SupervisorDecision,
)
from app.domain.errors import WorkspaceConflict
from app.main import app
from app.model_client import ModelUnavailable
from app.observability import operations
from app.services.contract_templates import ContractTemplateCatalog

SOURCE = "采购方：甲方名称为华明公司。\n供货方：第一次报价总金额9000元。\n双方最终确认总金额8000元，含税。\n签字由线下完成。"


class DeterministicAgentTeam:
    """测试夹具或替身：DeterministicAgentTeam，用于隔离验证本模块行为。"""

    def __init__(self, _settings):
        """初始化当前测试夹具或模拟角色的内部状态。"""
        self.analysis = SimpleNamespace(analyze=self.analyze)
        self.extraction = SimpleNamespace(extract=self.extract)
        self.review = SimpleNamespace(review=self.review_fields)
        self.coverage = SimpleNamespace(review=self.review_coverage)
        self.supervisor = SimpleNamespace(decide=self.decide)

    async def decide(self, phase, source_lines, observations, allowed_actions):
        """Exercise decide within this isolated test scenario."""
        action = "handoff" if "handoff" in allowed_actions else allowed_actions[0]
        return SupervisorDecision(action=action, decision_summary="离线测试：执行必要步骤后交用户审核")

    async def review_coverage(self, *_args):
        """Exercise review coverage within this isolated test scenario."""
        return CoverageResult()

    async def __aenter__(self):
        """Exercise aenter within this isolated test scenario."""
        return self

    async def __aexit__(self, *_):
        """Exercise aexit within this isolated test scenario."""
        pass

    async def analyze(self, _lines):
        """Exercise analyze within this isolated test scenario."""
        return NegotiationAnalysis(
            summary="双方协商设备买卖，最终金额为8000元，乙方名称待补充。",
            contract_type="设备买卖",
            search_query="设备买卖合同",
            key_terms=["设备", "买卖"],
            unresolved_points=["乙方名称未出现"],
        )

    async def extract(self, _lines, template_fields):
        """Exercise extract within this isolated test scenario."""
        candidates = {
            "甲方名称": {
                "value": "华明公司",
                "status": "found",
                "evidence": [{"line_number": 1, "quote": "甲方名称为华明公司"}],
            },
            "金额": {
                "value": "9000元",
                "status": "found",
                "evidence": [{"line_number": 2, "quote": "第一次报价总金额9000元"}],
                "alternatives": ["8000元", "9000元"],
            },
        }
        return ExtractionResult(
            fields=[
                {
                    "field_name": field["name"],
                    **candidates.get(field["name"], {"value": "", "status": "missing"}),
                }
                for field in template_fields
            ]
        )

    async def review_fields(self, _lines, template_fields, _extracted):
        """Exercise review fields within this isolated test scenario."""
        return ReviewResult(
            findings=[
                {
                    "field_name": "金额",
                    "severity": "conflict",
                    "issue": "9000元为初次报价；第3行最终确认8000元，请人工核对",
                }
            ]
            if any(field["name"] == "金额" for field in template_fields)
            else []
        )


class ContractWorkflowTests(TestCase):
    """测试夹具或替身：ContractWorkflowTests，用于隔离验证本模块行为。"""

    def setUp(self):
        # A developer's real .env must never enable paid API calls in tests.
        """准备隔离测试资源，避免访问生产数据库和真实模型。"""
        configuration_patch = patch.dict(Settings.model_config, {"env_file": None})
        configuration_patch.start()
        self.addCleanup(configuration_patch.stop)
        self.temporary = TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.resource_root = self.root / "contracts"
        self.template_folder = self.resource_root / "templates" / "Ministerial"
        self.template_folder.mkdir(parents=True)
        for name in ("设备买卖合同", "设备采购补充合同"):
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("甲方：{")
            paragraph.add_run("甲方名称}，金额：{金额}")
            document.add_table(rows=1, cols=1).cell(0, 0).text = "乙方：{乙方名称}，签字：{签字}"
            document.save(self.template_folder / f"{name}.docx")
        self.environment = patch.dict(
            os.environ,
            {
                "LAW_DATABASE_PATH": str(self.root / "test.db"),
                "LAW_RESOURCES_DIR": str(self.resource_root),
                "LAW_GENERATED_DIR": str(self.root / "generated"),
                "LAW_LLM_API_KEY": "",
                "LAW_LLM_BASE_URL": "",
                "LAW_LLM_MODEL": "",
            },
        )
        self.environment.start()
        get_settings.cache_clear()
        get_database.cache_clear()
        get_template_catalog.cache_clear()
        self.team_patch = patch("app.services.contract_workflow.create_agent_team", DeterministicAgentTeam)
        self.team_patch.start()
        self.client = TestClient(app)
        self.client.__enter__()
        self.headers = self.register_and_login("alice")

    def tearDown(self):
        """释放本测试创建的客户端、临时存储及配置替身。"""
        self.client.__exit__(None, None, None)
        self.team_patch.stop()
        self.environment.stop()
        get_settings.cache_clear()
        get_database.cache_clear()
        get_template_catalog.cache_clear()
        self.temporary.cleanup()

    def register_and_login(self, username):
        """Exercise register and login within this isolated test scenario."""
        response = self.client.post(
            "/api/v1/auth/register",
            json={"email": f"{username}@example.org", "password": "safe-password-123"},
        )
        self.assertEqual(response.status_code, 201, response.text)
        login = self.client.post(
            "/api/v1/auth/login", json={"email": f"{username}@example.org", "password": "safe-password-123"}
        )
        return {"Authorization": f"Bearer {login.json()['access_token']}"}

    def create_workspace(self):
        """Exercise create workspace within this isolated test scenario."""
        response = self.client.post(
            "/api/v1/workspaces", headers=self.headers, json={"title": "设备采购", "source_text": SOURCE}
        )
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def transition(self, workspace, action, payload=None, method="post", expected=200):
        """Exercise transition within this isolated test scenario."""
        body = {"revision": workspace["revision"], **(payload or {})}
        response = getattr(self.client, method)(
            f"/api/v1/workspaces/{workspace['id']}/{action}", headers=self.headers, json=body
        )
        self.assertEqual(response.status_code, expected, response.text)
        return response.json()

    def selected_workspace(self):
        """Exercise selected workspace within this isolated test scenario."""
        workspace = self.create_workspace()
        templates = self.client.get("/api/v1/contracts/templates", headers=self.headers).json()["items"]
        return self.transition(workspace, "template", {"template_id": templates[0]["id"]})

    def reviewed_workspace(self):
        """Exercise reviewed workspace within this isolated test scenario."""
        workspace = self.transition(self.selected_workspace(), "extract")
        decisions = [
            {"name": "甲方名称", "value": "华明公司", "action": "confirm"},
            {"name": "金额", "value": "8000元", "action": "confirm"},
            {"name": "乙方名称", "value": "新星公司", "action": "confirm", "note": "用户补充全称"},
            {"name": "签字", "value": "", "action": "omit", "note": "签署时线下手写"},
        ]
        return self.transition(workspace, "fields", {"decisions": decisions}, method="patch")

    def test_shared_omission_reason_allows_draft_and_later_completion(self):
        """Verify shared omission reason allows draft and later completion."""
        workspace = self.selected_workspace()
        decisions = [
            {"name": field["name"], "action": "omit", "note": "首轮纪要未提供，暂留空"}
            for field in workspace["fields"]
        ]
        workspace = self.transition(workspace, "fields", {"decisions": decisions}, method="patch")
        self.assertEqual(workspace["stage"], "field_review")
        workspace = self.transition(
            workspace,
            "fields",
            {"decisions": [{"name": "甲方名称", "value": "华明公司", "action": "confirm"}]},
            method="patch",
        )
        self.assertEqual(workspace["stage"], "ready_to_generate")
        preview = self.client.get(
            f"/api/v1/workspaces/{workspace['id']}/preview", headers=self.headers
        ).json()
        self.assertEqual(preview["pending_count"], 0)
        self.assertEqual(len(preview["omitted_fields"]), 3)
        workspace = self.transition(workspace, "generate", {"preview_acknowledged": True})
        old_artifact = workspace["artifacts"][0]["id"]
        workspace = self.transition(
            workspace,
            "fields",
            {"decisions": [{"name": "乙方名称", "value": "新星公司", "action": "confirm"}]},
            method="patch",
        )
        self.assertEqual(workspace["stage"], "ready_to_generate")
        self.assertEqual(
            next(field for field in workspace["fields"] if field["name"] == "乙方名称")["value"], "新星公司"
        )
        downloaded = self.client.get(f"/api/v1/contracts/{old_artifact}/download", headers=self.headers)
        self.assertNotIn("新星公司", Document(BytesIO(downloaded.content)).tables[0].cell(0, 0).text)

    def test_old_mapping_requires_reselection_without_relabeling_old_fields(self):
        """Verify old mapping requires reselection without relabeling old fields."""
        workspace = self.selected_workspace()
        with patch("app.documents.template_slots.SLOT_ENGINE_VERSION", "unused"):
            # The catalog owns the imported version used in its fingerprint.
            with patch("app.services.contract_templates.SLOT_ENGINE_VERSION", "new-test-version"):
                loaded = self.client.get(f"/api/v1/workspaces/{workspace['id']}", headers=self.headers).json()
                self.assertTrue(loaded["selected_template"]["requires_reselection"])
                self.assertEqual(loaded["fields"], workspace["fields"])
                self.transition(
                    workspace,
                    "fields",
                    {"decisions": [{"name": "甲方名称", "value": "公司", "action": "confirm"}]},
                    method="patch",
                    expected=409,
                )

    def test_coverage_reports_only_verifiable_gaps_and_keeps_extraction_on_failure(self):
        """Verify coverage reports only verifiable gaps and keeps extraction on failure."""

        class GappedTeam(DeterministicAgentTeam):
            """测试夹具或替身：GappedTeam，用于隔离验证本模块行为。"""

            async def review_coverage(self, *_args):
                """Exercise review coverage within this isolated test scenario."""
                return CoverageResult(
                    gaps=[
                        {
                            "summary": "最终金额需补充",
                            "evidence": {"line_number": 3, "quote": "最终确认总金额8000元"},
                            "suggestion": "核对金额字段",
                        },
                        {
                            "summary": "虚假遗漏",
                            "evidence": {"line_number": 1, "quote": "伪造原文"},
                            "suggestion": "不得保留",
                        },
                    ]
                )

        with patch("app.services.contract_workflow.create_agent_team", GappedTeam):
            workspace = self.transition(self.selected_workspace(), "extract")
        self.assertEqual(len(workspace["coverage_gaps"]), 1)
        self.assertEqual(workspace["coverage_status"], "unverified")

        class FailedCoverageTeam(DeterministicAgentTeam):
            """测试夹具或替身：FailedCoverageTeam，用于隔离验证本模块行为。"""

            async def review_coverage(self, *_args):
                """Exercise review coverage within this isolated test scenario."""
                raise ModelUnavailable("模拟覆盖服务失败")

        with patch("app.services.contract_workflow.create_agent_team", FailedCoverageTeam):
            workspace = self.transition(self.selected_workspace(), "extract")
        self.assertEqual(workspace["coverage_status"], "failed")
        self.assertEqual(
            next(field for field in workspace["fields"] if field["name"] == "甲方名称")["value"], "华明公司"
        )

    def test_typed_date_errors_are_flagged_and_cannot_be_confirmed(self):
        """Verify typed date errors are flagged and cannot be confirmed."""
        document = Document()
        document.add_paragraph("期限自____年____月____日")
        document.save(self.template_folder / "日期分格合同.docx")
        get_template_catalog().refresh()
        templates = self.client.get("/api/v1/contracts/templates", headers=self.headers).json()["items"]
        template = next(item for item in templates if item["name"] == "日期分格合同")
        workspace = self.transition(self.create_workspace(), "template", {"template_id": template["id"]})

        class BadDateTeam(DeterministicAgentTeam):
            """测试夹具或替身：BadDateTeam，用于隔离验证本模块行为。"""

            async def extract(self, _lines, template_fields):
                """Exercise extract within this isolated test scenario."""
                return ExtractionResult(
                    fields=[
                        {
                            "field_name": field["name"],
                            "value": "2026年11月1日",
                            "status": "found",
                            "evidence": [{"line_number": 1, "quote": "华明公司"}],
                        }
                        for field in template_fields
                    ]
                )

        with patch("app.services.contract_workflow.create_agent_team", BadDateTeam):
            workspace = self.transition(workspace, "extract")
        self.assertTrue(
            all(field["status"] == "unverified" and field["review_issues"] for field in workspace["fields"])
        )
        self.transition(
            workspace,
            "fields",
            {
                "decisions": [
                    {"name": field["name"], "value": field["value"], "action": "confirm"}
                    for field in workspace["fields"]
                ]
            },
            method="patch",
            expected=400,
        )
        loaded = self.client.get(f"/api/v1/workspaces/{workspace['id']}", headers=self.headers).json()
        self.assertEqual(loaded["revision"], workspace["revision"])

    def test_full_record_to_reviewed_docx_workflow(self):
        """Verify full record to reviewed docx workflow."""
        workspace = self.transition(self.create_workspace(), "analyze")
        self.assertTrue(workspace["recommendations"])
        recommendation = workspace["recommendations"][0]
        self.assertIn("match_reasons", recommendation)
        workspace = self.transition(workspace, "template", {"template_id": recommendation["id"]})
        workspace = self.transition(workspace, "extract")
        fields = {field["name"]: field for field in workspace["fields"]}
        self.assertEqual(fields["甲方名称"]["status"], "extracted")
        self.assertEqual(fields["金额"]["status"], "conflict")
        self.assertEqual(fields["乙方名称"]["status"], "missing")
        self.transition(workspace, "generate", {"preview_acknowledged": True}, expected=400)
        decisions = [
            {"name": name, "value": value, "action": "confirm"}
            for name, value in [("甲方名称", "华明公司"), ("金额", "8000元"), ("乙方名称", "新星公司")]
        ]
        decisions.append({"name": "签字", "action": "omit", "note": "线下签署"})
        workspace = self.transition(workspace, "fields", {"decisions": decisions}, method="patch")
        self.assertEqual(workspace["stage"], "ready_to_generate")
        amount = next(field for field in workspace["fields"] if field["name"] == "金额")
        self.assertEqual(amount["extracted_value"], "9000元")
        self.assertEqual(amount["value_origin"], "user_edit")
        preview = self.client.get(
            f"/api/v1/workspaces/{workspace['id']}/preview", headers=self.headers
        ).json()
        self.assertIn("8000元", preview["text"])
        self.assertIn("{签字}", preview["text"])
        self.transition(workspace, "generate", {"preview_acknowledged": False}, expected=400)
        workspace = self.transition(workspace, "generate", {"preview_acknowledged": True})
        artifact = workspace["artifacts"][0]
        downloaded = self.client.get(f"/api/v1/contracts/{artifact['id']}/download", headers=self.headers)
        self.assertEqual(downloaded.status_code, 200)
        document = Document(BytesIO(downloaded.content))
        self.assertIn("金额：8000元", document.paragraphs[0].text)
        self.assertIn("新星公司", document.tables[0].cell(0, 0).text)
        roles = {event["role"] for event in workspace["events"]}
        self.assertTrue(
            {"analysis_agent", "extraction_agent", "review_agent", "user", "renderer"}.issubset(roles)
        )

    def test_durable_runs_survive_telemetry_reset_and_enforce_ownership(self):
        """Verify durable runs survive telemetry reset and enforce ownership."""
        workspace = self.transition(self.create_workspace(), "analyze")
        url = f"/api/v1/workspaces/{workspace['id']}"
        runs = self.client.get(url + "/runs", headers=self.headers).json()["items"]
        self.assertEqual(runs[0]["status"], "completed")
        self.assertEqual(
            [step["name"] for step in runs[0]["steps"]],
            [
                "plan_next_action",
                "analyze_record",
                "plan_next_action",
                "retrieve_templates",
                "plan_next_action",
            ],
        )
        with operations.lock:
            operations.records.pop(workspace["id"], None)
        activity = self.client.get(url + "/activity", headers=self.headers).json()
        self.assertEqual(activity["harness"]["run_id"], runs[0]["run_id"])
        other = self.register_and_login("other-harness-user")
        self.assertEqual(self.client.get(url + "/runs", headers=other).status_code, 404)

    def test_failed_analysis_is_durable_without_overwriting_workspace(self):
        """Verify failed analysis is durable without overwriting workspace."""

        class FailingTeam(DeterministicAgentTeam):
            """测试夹具或替身：FailingTeam，用于隔离验证本模块行为。"""

            async def analyze(self, _lines):
                """Exercise analyze within this isolated test scenario."""
                raise ModelUnavailable("test model unavailable")

        workspace = self.create_workspace()
        with patch("app.services.contract_workflow.create_agent_team", FailingTeam):
            self.transition(workspace, "analyze", expected=503)
        runs = self.client.get(f"/api/v1/workspaces/{workspace['id']}/runs", headers=self.headers).json()[
            "items"
        ]
        self.assertEqual(runs[0]["status"], "failed")
        self.assertEqual(
            next(step for step in runs[0]["steps"] if step["name"] == "analyze_record")["status"], "failed"
        )
        saved = self.client.get(f"/api/v1/workspaces/{workspace['id']}", headers=self.headers).json()
        self.assertEqual(saved["revision"], workspace["revision"])

    def test_supervisor_repair_reaches_docx_only_after_explicit_human_approval(self):
        """Verify supervisor repair reaches docx only after explicit human approval."""

        class RepairTeam(DeterministicAgentTeam):
            """测试夹具或替身：RepairTeam，用于隔离验证本模块行为。"""

            def __init__(self, settings):
                """初始化当前测试夹具或模拟角色的内部状态。"""
                super().__init__(settings)
                self.extraction.repair = self.repair

            async def decide(self, phase, source, observations, allowed):
                """Exercise decide within this isolated test scenario."""
                if phase == "field_review" and "repair_fields" in allowed:
                    return SupervisorDecision(
                        action="repair_fields",
                        decision_summary="修正已作废报价",
                        field_ids=[observations["repairable_fields"][0]["name"]],
                    )
                return await super().decide(phase, source, observations, allowed)

            async def repair(self, source, specs, previous):
                """Exercise repair within this isolated test scenario."""
                return ExtractionResult(
                    fields=[
                        {
                            "field_name": "金额",
                            "value": "8000元",
                            "status": "found",
                            "evidence": [{"line_number": 3, "quote": "最终确认总金额8000元"}],
                        }
                    ]
                )

            async def review_fields(self, lines, specs, fields):
                """Exercise review fields within this isolated test scenario."""
                if any(field["name"] == "金额" and field["value"] == "8000元" for field in fields):
                    return ReviewResult()
                return await super().review_fields(lines, specs, fields)

        workspace = self.selected_workspace()
        with patch("app.services.contract_workflow.create_agent_team", RepairTeam):
            workspace = self.transition(workspace, "extract")
        amount = next(field for field in workspace["fields"] if field["name"] == "金额")
        self.assertEqual(amount["value"], "8000元")
        self.assertEqual(amount["status"], "extracted")
        self.assertTrue(any(event["role"] == "supervisor_agent" for event in workspace["events"]))
        self.transition(workspace, "generate", {"preview_acknowledged": True}, expected=400)
        decisions = [
            {"name": "甲方名称", "value": "华明公司", "action": "confirm"},
            {"name": "金额", "value": "8000元", "action": "confirm"},
            {"name": "乙方名称", "action": "omit", "note": "待补充"},
            {"name": "签字", "action": "omit", "note": "线下签署"},
        ]
        workspace = self.transition(workspace, "fields", {"decisions": decisions}, method="patch")
        generated = self.transition(workspace, "generate", {"preview_acknowledged": True})
        output = Document(self.root / "generated" / generated["artifacts"][0]["filename"])
        self.assertIn("金额：8000元", output.paragraphs[0].text)

    def test_literal_requested_title_survives_generic_model_search_query(self):
        """Verify literal requested title survives generic model search query."""
        response = self.client.post(
            "/api/v1/workspaces",
            headers=self.headers,
            json={"title": "明确合同", "source_text": "双方确认：需要设备采购补充合同。"},
        )
        self.assertEqual(response.status_code, 201)
        analyzed = self.transition(response.json(), "analyze")
        self.assertEqual(analyzed["recommendations"][0]["name"], "设备采购补充合同")
        self.assertTrue(analyzed["recommendations"][0]["explicitly_requested"])

    def test_email_only_authentication_and_activity_ownership(self):
        """Verify email only authentication and activity ownership."""
        workspace = self.transition(self.create_workspace(), "analyze")
        self.assertIn("last_run", workspace)
        activity = self.client.get(f"/api/v1/workspaces/{workspace['id']}/activity", headers=self.headers)
        self.assertEqual(activity.status_code, 200)
        self.assertEqual(activity.json()["status"], "completed")
        self.assertEqual(activity.json()["requests_started"], 0)  # deterministic test team is not a real LLM
        other_headers = self.register_and_login("another-user")
        hidden = self.client.get(f"/api/v1/workspaces/{workspace['id']}/activity", headers=other_headers)
        self.assertEqual(hidden.status_code, 404)
        login = self.client.post(
            "/api/v1/auth/login", json={"email": "ALICE@example.org", "password": "safe-password-123"}
        )
        self.assertEqual(login.status_code, 200)
        self.assertNotIn("username", login.json()["user"])
        duplicate = self.client.post(
            "/api/v1/auth/register", json={"email": "ALICE@example.org", "password": "safe-password-123"}
        )
        self.assertEqual(duplicate.status_code, 409)

    def test_preview_reports_pending_work_and_structured_table_blocks(self):
        """Verify preview reports pending work and structured table blocks."""
        workspace = self.selected_workspace()
        preview = self.client.get(
            f"/api/v1/workspaces/{workspace['id']}/preview", headers=self.headers
        ).json()
        self.assertEqual(preview["pending_count"], len(workspace["fields"]))
        self.assertIn("table", [block["type"] for block in preview["blocks"]])
        refreshed = self.client.get(f"/api/v1/workspaces/{workspace['id']}", headers=self.headers).json()
        self.assertTrue(all(field["display_name"] for field in refreshed["fields"]))

    def test_activity_can_be_polled_while_analysis_request_is_still_pending(self):
        """Verify activity can be polled while analysis request is still pending."""
        started, release = Event(), Event()

        class DelayedTeam(DeterministicAgentTeam):
            """测试夹具或替身：DelayedTeam，用于隔离验证本模块行为。"""

            async def analyze(self, lines):
                """Exercise analyze within this isolated test scenario."""
                started.set()
                while not release.is_set():
                    await asyncio.sleep(0.01)
                return await super().analyze(lines)

        workspace = self.create_workspace()
        with (
            patch("app.services.contract_workflow.create_agent_team", DelayedTeam),
            ThreadPoolExecutor(max_workers=1) as executor,
        ):
            result = executor.submit(self.transition, workspace, "analyze")
            try:
                self.assertTrue(started.wait(timeout=5))
                live = self.client.get(f"/api/v1/workspaces/{workspace['id']}/activity", headers=self.headers)
                self.assertEqual(live.status_code, 200)
                self.assertEqual(live.json()["status"], "running")
                self.assertEqual(live.json()["requests_started"], 0)
            finally:
                release.set()
            self.assertEqual(result.result(timeout=5)["last_run"]["status"], "completed")

    def test_completion_is_not_published_before_database_commit(self):
        """Verify completion is not published before database commit."""
        workspace = self.create_workspace()

        def reject_commit(*_args, **_kwargs):
            """Exercise reject commit within this isolated test scenario."""
            self.assertEqual(operations.snapshot(workspace["id"])["status"], "running")
            raise WorkspaceConflict("模拟版本冲突")

        with patch(
            "app.repositories.contract_workspaces.ContractWorkspaceRepository.save", side_effect=reject_commit
        ):
            self.transition(workspace, "analyze", expected=409)
        activity = self.client.get(
            f"/api/v1/workspaces/{workspace['id']}/activity", headers=self.headers
        ).json()
        self.assertEqual(activity["status"], "failed")

    def test_manual_template_choice_is_durable_and_never_overridden(self):
        """Verify manual template choice is durable and never overridden."""
        workspace = self.transition(self.create_workspace(), "analyze")
        items = self.client.get("/api/v1/contracts/templates?limit=1&offset=1", headers=self.headers).json()
        self.assertEqual(items["total"], 2)
        selected_id = items["items"][0]["id"]
        workspace = self.transition(workspace, "template", {"template_id": selected_id})
        self.transition(workspace, "analyze", expected=400)
        restored = self.client.get(f"/api/v1/workspaces/{workspace['id']}", headers=self.headers).json()
        self.assertEqual(restored["selected_template"]["id"], selected_id)

    def test_stale_revisions_and_user_ownership(self):
        """Verify stale revisions and user ownership."""
        workspace = self.selected_workspace()
        old_revision = workspace["revision"]
        updated = self.transition(workspace, "extract")
        self.transition(
            workspace,
            "fields",
            {"decisions": [{"name": "金额", "value": "1", "action": "confirm"}]},
            method="patch",
            expected=409,
        )
        self.assertGreater(updated["revision"], old_revision)
        other_headers = self.register_and_login("bob")
        self.assertEqual(
            self.client.get(f"/api/v1/workspaces/{workspace['id']}", headers=other_headers).status_code, 404
        )
        self.assertEqual(
            self.client.post(
                f"/api/v1/workspaces/{workspace['id']}/extract",
                headers=other_headers,
                json={"revision": updated["revision"]},
            ).status_code,
            404,
        )

    def test_user_edits_require_explicit_reextraction_and_unknown_fields_are_rejected(self):
        """Verify user edits require explicit reextraction and unknown fields are rejected."""
        workspace = self.reviewed_workspace()
        self.transition(workspace, "extract", expected=409)
        self.transition(
            workspace,
            "fields",
            {"decisions": [{"name": "伪造字段", "value": "x"}]},
            method="patch",
            expected=400,
        )
        reset = self.transition(workspace, "extract", {"replace_reviewed": True})
        self.assertEqual(reset["stage"], "field_review")
        self.assertFalse(any(field["status"] == "confirmed" for field in reset["fields"]))

    def test_omitting_a_field_requires_a_reason(self):
        """Verify omitting a field requires a reason."""
        workspace = self.transition(self.selected_workspace(), "extract")
        self.transition(
            workspace,
            "fields",
            {"decisions": [{"name": "签字", "action": "omit"}]},
            method="patch",
            expected=400,
        )

    def test_all_failed_batches_skip_coverage_and_keep_public_recovery_details(self):
        """全部提取失败时保留任务和错误分类，并跳过后续模型检查。"""

        class FailedTeam(DeterministicAgentTeam):
            """模拟提取格式失败。"""

            async def extract(self, *_):
                """返回可公开的结构校验错误。"""
                raise ModelUnavailable("模型结果格式无效", code="model_schema", hint="重试失败批次")

            async def review_coverage(self, *_):
                """禁止在全部批次失败后请求覆盖检查。"""
                raise AssertionError("coverage should be skipped")

        workspace = self.transition(self.create_workspace(), "analyze")
        workspace = self.transition(
            workspace, "template", {"template_id": workspace["recommendations"][0]["id"]}
        )
        with patch("app.services.contract_workflow.create_agent_team", FailedTeam):
            workspace = self.transition(workspace, "extract")
        self.assertEqual(workspace["coverage_status"], "skipped")
        self.assertEqual(workspace["extraction_failures"][0]["error"]["code"], "model_schema")
        self.assertEqual(workspace["last_run"]["status"], "partial")
        failed_step = next(
            step for step in workspace["last_run"]["harness"]["steps"] if step["status"] == "failed"
        )
        self.assertEqual(failed_step["error"]["code"], "model_schema")

    def test_failed_batches_can_retry_without_overwriting_confirmed_fields(self):
        """Verify failed batches can retry without overwriting confirmed fields."""
        document = Document()
        document.add_paragraph(" ".join("{字段" + str(index) + "}" for index in range(10)))
        document.save(self.template_folder / "分批测试合同.docx")
        get_template_catalog().refresh()
        get_settings().fields_per_batch = 5
        template_id = next(
            item["id"] for item in get_template_catalog().list_templates() if item["name"] == "分批测试合同"
        )
        workspace = self.transition(self.create_workspace(), "template", {"template_id": template_id})

        class PartialFailureTeam(DeterministicAgentTeam):
            """测试夹具或替身：PartialFailureTeam，用于隔离验证本模块行为。"""

            attempted_batches = []

            async def extract(self, lines, fields):
                """Exercise extract within this isolated test scenario."""
                names = [field["name"] for field in fields]
                self.attempted_batches.append(names)
                if "字段5" in names and len(self.attempted_batches) <= 2:
                    raise ModelUnavailable("模拟一批限流")
                return await super().extract(lines, fields)

        with patch("app.services.contract_workflow.create_agent_team", PartialFailureTeam):
            workspace = self.transition(workspace, "extract")
            self.assertEqual(len(workspace["extraction_failures"]), 1)
            workspace = self.transition(
                workspace,
                "fields",
                {"decisions": [{"name": "字段0", "value": "人工值", "action": "confirm"}]},
                method="patch",
            )
            workspace = self.transition(workspace, "extract", {"retry_failed": True})
        self.assertEqual(workspace["extraction_failures"], [])
        self.assertEqual(
            PartialFailureTeam.attempted_batches[-1], ["字段5", "字段6", "字段7", "字段8", "字段9"]
        )
        confirmed = next(field for field in workspace["fields"] if field["name"] == "字段0")
        self.assertEqual((confirmed["status"], confirmed["value"]), ("confirmed", "人工值"))

    def test_manual_review_and_generation_work_without_model_configuration(self):
        """Verify manual review and generation work without model configuration."""
        workspace = self.selected_workspace()
        self.assertEqual(len(workspace["fields"]), 4)
        decisions = [
            {"name": field["name"], "action": "confirm", "value": "人工填写"} for field in workspace["fields"]
        ]
        workspace = self.transition(workspace, "fields", {"decisions": decisions}, method="patch")
        generated = self.transition(workspace, "generate", {"preview_acknowledged": True})
        self.assertEqual(len(generated["artifacts"]), 1)

    def test_forged_evidence_remains_unverified(self):
        """Verify forged evidence remains unverified."""

        class ForgedEvidenceTeam(DeterministicAgentTeam):
            """测试夹具或替身：ForgedEvidenceTeam，用于隔离验证本模块行为。"""

            async def extract(self, _lines, template_fields):
                """Exercise extract within this isolated test scenario."""
                return ExtractionResult(
                    fields=[
                        {
                            "field_name": field["name"],
                            "value": "虚构",
                            "status": "found",
                            "evidence": [{"line_number": 99, "quote": "不存在的原文"}],
                        }
                        for field in template_fields
                    ]
                )

        with patch("app.services.contract_workflow.create_agent_team", ForgedEvidenceTeam):
            workspace = self.transition(self.selected_workspace(), "extract")
        field = next(field for field in workspace["fields"] if field["name"] == "甲方名称")
        self.assertEqual(field["status"], "unverified")
        self.assertEqual(field["evidence"], [])

    def test_template_file_change_invalidates_selected_version(self):
        """Verify template file change invalidates selected version."""
        workspace = self.selected_workspace()
        path = get_template_catalog().get_template(workspace["selected_template"]["id"]).path
        document = Document(path)
        document.add_paragraph("新增：{新字段}")
        document.save(path)
        self.transition(workspace, "extract", expected=409)

    def test_engine_upgrade_invalidates_old_field_locations(self):
        """Verify engine upgrade invalidates old field locations."""
        workspace = self.selected_workspace()
        with patch("app.services.contract_templates.SLOT_ENGINE_VERSION", "next-engine"):
            self.transition(workspace, "extract", expected=409)
            response = self.client.get(f"/api/v1/workspaces/{workspace['id']}/preview", headers=self.headers)
            self.assertEqual(response.status_code, 409)

    def test_original_template_manual_api_workflow(self):
        """Verify original template manual api workflow."""
        for path in self.template_folder.glob("*.docx"):
            document = Document()
            document.add_paragraph("采购人____")
            table = document.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "货物名称"
            document.save(path)
        workspace = self.selected_workspace()
        self.assertEqual(len(workspace["fields"]), 2)
        decisions = [
            {"name": field["name"], "value": "华明公司" if index == 0 else "设备A", "action": "confirm"}
            for index, field in enumerate(workspace["fields"])
        ]
        workspace = self.transition(workspace, "fields", {"decisions": decisions}, method="patch")
        preview = self.client.get(
            f"/api/v1/workspaces/{workspace['id']}/preview", headers=self.headers
        ).json()
        self.assertIn("华明公司", preview["text"])
        generated = self.transition(workspace, "generate", {"preview_acknowledged": True})
        output = Document(self.root / "generated" / generated["artifacts"][0]["filename"])
        self.assertEqual(output.tables[0].cell(1, 0).text, "设备A")

    def test_visual_content_requires_separate_user_approval_before_export(self):
        """Verify visual content requires separate user approval before export."""
        for path in self.template_folder.glob("*.docx"):
            document = Document()
            document.add_paragraph("采购人____").add_run()._r.append(
                parse_xml(
                    '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                    'xmlns:v="urn:schemas-microsoft-com:vml"><v:shape id="retained-graphic"/></w:pict>'
                )
            )
            document.save(path)
        workspace = self.selected_workspace()
        self.assertTrue(workspace["selected_template"]["requires_visual_review"])
        workspace = self.transition(
            workspace,
            "fields",
            {
                "decisions": [
                    {"name": workspace["fields"][0]["name"], "value": "华明公司", "action": "confirm"}
                ]
            },
            method="patch",
        )
        self.transition(workspace, "generate", {"preview_acknowledged": True}, expected=400)
        generated = self.transition(
            workspace, "generate", {"preview_acknowledged": True, "visual_review_acknowledged": True}
        )
        self.assertTrue(generated["artifacts"][0]["visual_review_acknowledged"])
        output = Document(self.root / "generated" / generated["artifacts"][0]["filename"])
        self.assertEqual(len(output.element.xpath(".//w:pict")), 1)

    def test_model_failure_keeps_imported_record(self):
        """Verify model failure keeps imported record."""
        workspace = self.create_workspace()
        with patch("app.services.contract_workflow.create_agent_team", real_agent_team):
            self.transition(workspace, "analyze", expected=503)
        restored = self.client.get(f"/api/v1/workspaces/{workspace['id']}", headers=self.headers).json()
        self.assertEqual(restored["source_text"], SOURCE)
        self.assertEqual(restored["revision"], workspace["revision"])

    def test_import_and_template_download(self):
        """Verify import and template download."""
        uploaded = self.client.post(
            "/api/v1/sources/parse",
            headers=self.headers,
            files={"file": ("纪要.txt", SOURCE.encode(), "text/plain")},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        self.assertEqual(uploaded.json()["source_text"], SOURCE)
        self.assertEqual(
            self.client.post(
                "/api/v1/sources/parse",
                headers=self.headers,
                files={"file": ("record.exe", b"not supported")},
            ).status_code,
            400,
        )
        document = Document()
        document.add_paragraph("首轮报价9000元")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "最终商定8000元"
        document.add_paragraph("会议结束并确认上述内容")
        stream = BytesIO()
        document.save(stream)
        imported = self.client.post(
            "/api/v1/sources/parse", headers=self.headers, files={"file": ("纪要.docx", stream.getvalue())}
        )
        self.assertEqual(imported.status_code, 200)
        imported_text = imported.json()["source_text"]
        self.assertLess(imported_text.index("首轮报价"), imported_text.index("最终商定"))
        self.assertLess(imported_text.index("最终商定"), imported_text.index("会议结束"))
        template_id = get_template_catalog().list_templates()[0]["id"]
        self.assertEqual(
            self.client.get(
                f"/api/v1/contracts/templates/{template_id}/download", headers=self.headers
            ).status_code,
            200,
        )

    def test_delete_removes_owned_files_and_revokes_access(self):
        """Verify delete removes owned files and revokes access."""
        workspace = self.transition(self.reviewed_workspace(), "generate", {"preview_acknowledged": True})
        artifact = workspace["artifacts"][0]
        self.assertTrue((self.root / "generated" / artifact["filename"]).exists())
        other_headers = self.register_and_login("bob")
        self.assertEqual(
            self.client.get(
                f"/api/v1/contracts/{artifact['id']}/download", headers=other_headers
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.delete(f"/api/v1/workspaces/{workspace['id']}", headers=self.headers).status_code, 200
        )
        self.assertFalse((self.root / "generated" / artifact["filename"]).exists())
        self.assertEqual(
            self.client.get(f"/api/v1/contracts/{artifact['id']}/download", headers=self.headers).status_code,
            404,
        )

    def test_catalog_uses_actual_docx_placeholders_and_preserves_split_runs(self):
        """Verify catalog uses actual docx placeholders and preserves split runs."""
        placeholder_folder = self.template_folder / "占位符"
        placeholder_folder.mkdir()
        (placeholder_folder / "设备买卖合同.json").write_text(
            json.dumps({"placeholders": ["过期字段"]}), encoding="utf-8"
        )
        catalog = ContractTemplateCatalog(self.resource_root, self.root / "generated")
        template = next(item for item in catalog.list_templates() if item["name"] == "设备买卖合同")
        details = catalog.describe_template(template["id"])
        self.assertIn("甲方名称", details["placeholders"])
        self.assertNotIn("过期字段", details["placeholders"])
        generated = catalog.generate_contract(template["id"], {"甲方名称": "华明公司", "金额": "8000元"})
        result = Document(self.root / "generated" / generated["filename"])
        self.assertEqual(result.paragraphs[0].text, "甲方：华明公司，金额：8000元")

    def test_auth_logout_and_safe_health_metadata(self):
        """Verify auth logout and safe health metadata."""
        health = self.client.get("/api/v1/health").json()
        self.assertFalse(health["model_configured"])
        self.assertIn("LAW_LLM_BASE_URL", health["missing_model_settings"])
        self.assertEqual(self.client.post("/api/v1/auth/logout", headers=self.headers).status_code, 200)
        self.assertEqual(self.client.get("/api/v1/workspaces", headers=self.headers).status_code, 401)

    def test_generated_snapshot_does_not_change_when_reviewed_values_change(self):
        """Verify generated snapshot does not change when reviewed values change."""
        reviewed = self.reviewed_workspace()
        generated = self.transition(reviewed, "generate", {"preview_acknowledged": True})
        artifact = generated["artifacts"][0]
        updated = self.transition(
            generated,
            "fields",
            {"decisions": [{"name": "金额", "value": "8500元", "action": "confirm"}]},
            method="patch",
        )
        self.transition(generated, "generate", {"preview_acknowledged": True}, expected=409)
        old_document = Document(self.root / "generated" / artifact["filename"])
        self.assertIn("8000元", old_document.paragraphs[0].text)
        preview = self.client.get(f"/api/v1/workspaces/{updated['id']}/preview", headers=self.headers).json()
        self.assertIn("8500元", preview["text"])

    def test_source_character_limit_is_enforced_without_model_calls(self):
        """Verify source character limit is enforced without model calls."""
        response = self.client.post(
            "/api/v1/workspaces",
            headers=self.headers,
            json={"source_text": "长" * (get_settings().max_source_chars + 1)},
        )
        self.assertEqual(response.status_code, 400)
