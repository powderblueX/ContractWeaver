"""Original-template detection, lossless blank restoration and safe rendering."""

from base64 import b64decode
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document
from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml

from app.documents.template_slots import compile_template_slots
from app.documents.word_documents import iter_document_paragraphs
from app.documents.word_integrity import visual_fingerprint
from app.documents.word_text import paragraph_text
from app.domain.errors import InvalidContractInput
from app.services.contract_templates import ContractTemplateCatalog


class OriginalTemplateTests(TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.folder = self.root / "templates" / "Ministerial"
        self.folder.mkdir(parents=True)
        self.path = self.folder / "原始合同.docx"
        self.catalog = ContractTemplateCatalog(self.root, self.root / "generated")

    def describe(self, document):
        document.save(self.path)
        self.catalog.refresh()
        identifier = self.catalog.list_templates()[0]["id"]
        return identifier, self.catalog.describe_template(identifier)

    def test_original_without_markers_extracts_table_and_split_run_blanks(self):
        document = Document()
        paragraph = document.add_paragraph("采购人：")
        paragraph.add_run("__").bold = True
        paragraph.add_run("__")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "产品名称"
        table.cell(0, 1).text = "数量"
        identifier, details = self.describe(document)
        original_bytes = self.path.read_bytes()
        self.assertTrue(details["fillable"])
        self.assertEqual(len(details["fields"]), 3)
        self.assertTrue(any("产品名称" in field["context"] for field in details["fields"]))
        values = {field["name"]: f"确认值{index}" for index, field in enumerate(details["fields"])}
        artifact = self.catalog.generate_contract(identifier, values)
        output = Document(self.root / "generated" / artifact["filename"])
        text = "\n".join(paragraph.text for paragraph in iter_document_paragraphs(output))
        self.assertTrue(all(value in text for value in values.values()))
        self.assertNotIn("·位置", text)
        self.assertEqual(self.path.read_bytes(), original_bytes)
        self.assertTrue(output.paragraphs[0].runs[1].bold)

    def test_adjacent_labels_and_repeated_amounts_are_separate(self):
        document = Document()
        document.add_paragraph("出卖人：签订地点：")
        document.add_paragraph("金额____，金额____")
        _, details = self.describe(document)
        self.assertEqual(len(details["fields"]), 4)
        self.assertEqual(len(set(details["placeholders"])), 4)

    def test_omitted_generated_slots_restore_original_blanks(self):
        document = Document()
        document.add_paragraph("甲方____，说明：")
        document.add_table(rows=1, cols=1)
        identifier, details = self.describe(document)
        first, *omitted = details["placeholders"]
        values = {first: "甲公司"}
        preview = self.catalog.preview_contract(identifier, values, omitted)
        self.assertNotIn("·位置", preview["text"])
        artifact = self.catalog.generate_contract(identifier, values, omitted)
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(output.paragraphs[0].text, "甲方甲公司，说明：")
        self.assertEqual(output.tables[0].cell(0, 0).text, "")
        self.assertEqual(artifact["missing_fields"], sorted(omitted))

    def test_underlined_whitespace_and_internal_spaces_are_detected(self):
        document = Document()
        paragraph = document.add_paragraph("日期")
        paragraph.add_run("   ").underline = True
        paragraph.add_run("年")
        document.add_paragraph("首期  元，次期  元")
        _, details = self.describe(document)
        self.assertEqual(len(details["fields"]), 3)

    def test_underlined_blank_lines_are_not_merged_across_a_real_line_break(self):
        document = Document()
        paragraph = document.add_paragraph("具体事宜：")
        run = paragraph.add_run("        ")
        run.underline = True
        run.add_break()
        run.add_text("        ")
        paragraph.add_run("。")
        identifier, details = self.describe(document)
        self.assertEqual(len(details["fields"]), 2)
        first, second = details["placeholders"]
        artifact = self.catalog.generate_contract(identifier, {first: "第一行", second: "第二行"})
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(output.paragraphs[0].text, "具体事宜：第一行\n第二行。")

    def test_unsupported_control_is_not_silently_filled(self):
        document = Document()
        document.add_paragraph("甲方____")
        document.element.body.append(OxmlElement("w:sdt"))
        identifier, details = self.describe(document)
        self.assertFalse(details["fillable"])
        self.assertIn("内容控件", details["unsupported_structures"])
        with self.assertRaises(InvalidContractInput):
            self.catalog.generate_contract(identifier, {details["placeholders"][0]: "甲公司"})

    def test_compilation_is_deterministic_and_explicit_markers_stay_supported(self):
        document = Document()
        document.add_paragraph("甲方：{甲方名称}，金额____")
        identifier, details = self.describe(document)
        self.assertEqual(details, self.catalog.describe_template(identifier))
        self.assertEqual(len(details["fields"]), 2)
        compiled = Document(self.path)
        first = compile_template_slots(compiled)
        second = compile_template_slots(compiled)
        self.assertEqual(first.generated_slots, 1)
        self.assertEqual(second.generated_slots, 0)

    def test_static_vml_separator_lines_are_preserved(self):
        document = Document()
        paragraph = document.add_paragraph("甲方____")
        paragraph.add_run()._r.append(
            parse_xml(
                '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:v="urn:schemas-microsoft-com:vml"><v:group><v:line from="0,0" to="10,0"/></v:group></w:pict>'
            )
        )
        identifier, details = self.describe(document)
        self.assertTrue(details["fillable"])
        artifact = self.catalog.generate_contract(identifier, {details["placeholders"][0]: "甲公司"})
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(len(output.element.xpath(".//w:pict")), 1)
        self.assertEqual(output.paragraphs[0].text, "甲方甲公司")

    def test_header_controls_are_blocked_too(self):
        document = Document()
        document.add_paragraph("甲方____")
        document.sections[0].header._element.append(OxmlElement("w:sdt"))
        _, details = self.describe(document)
        self.assertFalse(details["fillable"])

    def test_first_and_even_page_headers_are_detected_without_creating_new_parts(self):
        document = Document()
        document.add_paragraph("正文____")
        section = document.sections[0]
        section.first_page_header.paragraphs[0].text = "首页____"
        section.even_page_footer.paragraphs[0].text = "偶数页____"
        identifier, details = self.describe(document)
        self.assertEqual(len(details["fields"]), 3)
        values = {field["name"]: "确认" for field in details["fields"]}
        artifact = self.catalog.generate_contract(identifier, values)
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(output.sections[0].first_page_header.paragraphs[0].text, "首页确认")
        self.assertEqual(output.sections[0].even_page_footer.paragraphs[0].text, "偶数页确认")

    def test_separator_notes_are_not_mistaken_for_business_notes(self):
        document = Document()
        document.add_paragraph("正文____")
        xml = '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="-1" w:type="separator"/></w:footnotes>'
        part = Part(
            PackURI("/word/footnotes.xml"), CONTENT_TYPE.WML_FOOTNOTES, xml.encode(), document.part.package
        )
        document.part.relate_to(part, RELATIONSHIP_TYPE.FOOTNOTES)
        _, details = self.describe(document)
        self.assertTrue(details["fillable"])
        part._blob = xml.replace(
            "</w:footnotes>",
            '<w:footnote w:id="1"><w:p><w:r><w:t>实际脚注</w:t></w:r></w:p></w:footnote></w:footnotes>',
        ).encode()
        _, details = self.describe(document)
        self.assertFalse(details["fillable"])
        self.assertIn("脚注或尾注", details["unsupported_structures"])

    def test_merged_cell_continuations_are_not_invisible_input_slots(self):
        document = Document()
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).merge(table.cell(1, 0))
        _, details = self.describe(document)
        self.assertTrue(details["fillable"])
        self.assertEqual(len(details["fields"]), 1)

    def test_legacy_form_fields_are_blocked(self):
        document = Document()
        paragraph = document.add_paragraph("甲方____")
        field = OxmlElement("w:fldSimple")
        field.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr", " MERGEFIELD Party ")
        paragraph._p.append(field)
        _, details = self.describe(document)
        self.assertFalse(details["fillable"])
        self.assertIn("邮件合并或交互式填写域", details["unsupported_structures"])

    def test_mixed_graphic_runs_are_filled_without_removing_graphics(self):
        document = Document()
        paragraph = document.add_paragraph()
        run = paragraph.add_run("甲方____")
        run._r.append(
            parse_xml(
                '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:v="urn:schemas-microsoft-com:vml"><v:line from="0,0" to="10,0"/></w:pict>'
            )
        )
        identifier, details = self.describe(document)
        self.assertTrue(details["fillable"])
        artifact = self.catalog.generate_contract(identifier, {details["placeholders"][0]: "甲公司"})
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(len(output.element.xpath(".//w:pict")), 1)
        self.assertEqual(output.paragraphs[0].text, "甲方甲公司")

    def test_images_are_preserved_and_require_explicit_visual_review(self):
        image = b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
        )
        document = Document()
        paragraph = document.add_paragraph()
        run = paragraph.add_run("甲方____")
        run.add_picture(BytesIO(image))
        before = visual_fingerprint(document)
        identifier, details = self.describe(document)
        self.assertTrue(details["fillable"])
        self.assertTrue(details["requires_visual_review"])
        artifact = self.catalog.generate_contract(identifier, {details["placeholders"][0]: "甲公司"})
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(visual_fingerprint(output), before)
        self.assertEqual(output.paragraphs[0].text, "甲方甲公司")

    def test_word_text_box_content_is_editable_without_changing_its_frame(self):
        document = Document()
        paragraph = document.add_paragraph("买方____")
        paragraph.add_run()._r.append(
            parse_xml(
                '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:v="urn:schemas-microsoft-com:vml"><v:shape id="party-box" style="width:100pt;height:40pt">'
                "<v:textbox><w:txbxContent><w:p><w:r><w:t>卖方____</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>"
            )
        )
        before = visual_fingerprint(document)
        identifier, details = self.describe(document)
        self.assertTrue(details["fillable"])
        self.assertEqual(len(details["fields"]), 2)
        values = {field["name"]: f"公司{index}" for index, field in enumerate(details["fields"])}
        artifact = self.catalog.generate_contract(identifier, values)
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(visual_fingerprint(output), before)
        text = "\n".join(paragraph_text(item) for item in iter_document_paragraphs(output))
        self.assertIn("买方公司0", text)
        self.assertIn("卖方公司1", text)

    def test_empty_text_box_is_one_reviewable_candidate(self):
        document = Document()
        document.add_paragraph().add_run()._r.append(
            parse_xml(
                '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:v="urn:schemas-microsoft-com:vml"><v:shape><v:textbox><w:txbxContent>'
                "<w:p/><w:p/></w:txbxContent></v:textbox></v:shape></w:pict>"
            )
        )
        _, details = self.describe(document)
        self.assertTrue(details["fillable"])
        self.assertEqual(len(details["fields"]), 1)

    def test_hyperlink_text_can_be_filled_without_losing_its_relationship(self):
        document = Document()
        paragraph = document.add_paragraph()
        identifier = document.part.relate_to(
            "https://example.com", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        link = parse_xml(
            '<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            "<w:r><w:t>联系人____</w:t></w:r></w:hyperlink>"
        )
        link.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", identifier)
        paragraph._p.append(link)
        template_id, details = self.describe(document)
        artifact = self.catalog.generate_contract(template_id, {details["placeholders"][0]: "张先生"})
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(paragraph_text(output.paragraphs[0]), "联系人张先生")
        self.assertEqual(output.part.rels[identifier].target_ref, "https://example.com")

    def test_multiple_compatibility_text_box_branches_remain_blocked(self):
        document = Document()
        box = "<w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>甲方____</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>"
        document.add_paragraph().add_run()._r.append(
            parse_xml(
                '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
                'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:v="urn:schemas-microsoft-com:vml"><mc:Choice Requires="v">'
                + box
                + "</mc:Choice><mc:Fallback>"
                + box.replace("甲方", "乙方")
                + "</mc:Fallback></mc:AlternateContent>"
            )
        )
        _, details = self.describe(document)
        self.assertFalse(details["fillable"])
        self.assertTrue(any("兼容分支" in reason for reason in details["unsupported_structures"]))

    def test_equal_compatibility_branches_share_one_field_and_are_both_filled(self):
        document = Document()
        box = '<w:pict><v:shape id="same-box"><v:textbox><w:txbxContent><w:p><w:r><w:t>甲方____</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>'
        document.add_paragraph().add_run()._r.append(
            parse_xml(
                '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
                'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                'xmlns:v="urn:schemas-microsoft-com:vml"><mc:Choice Requires="v">'
                + box
                + "</mc:Choice><mc:Fallback>"
                + box
                + "</mc:Fallback></mc:AlternateContent>"
            )
        )
        before = visual_fingerprint(document)
        identifier, details = self.describe(document)
        self.assertTrue(details["fillable"])
        self.assertEqual(len(details["fields"]), 1)
        values = {details["placeholders"][0]: "甲公司"}
        self.assertEqual(self.catalog.preview_contract(identifier, values)["text"].count("甲公司"), 1)
        artifact = self.catalog.generate_contract(identifier, values)
        output = Document(self.root / "generated" / artifact["filename"])
        texts = [paragraph_text(item) for item in iter_document_paragraphs(output)]
        self.assertEqual(texts.count("甲方甲公司"), 2)
        self.assertEqual(visual_fingerprint(output), before)

    def test_colon_before_layout_breaks_does_not_delete_the_breaks(self):
        document = Document()
        document.add_paragraph("甲方：\n签字")
        document.add_paragraph("乙方：\t签字")
        identifier, details = self.describe(document)
        self.assertTrue(details["fillable"])
        self.assertEqual(len(details["fields"]), 2)
        values = {field["name"]: f"公司{index}" for index, field in enumerate(details["fields"])}
        artifact = self.catalog.generate_contract(identifier, values)
        output = Document(self.root / "generated" / artifact["filename"])
        self.assertEqual(paragraph_text(output.paragraphs[0]), "甲方：公司0\n签字")
        self.assertEqual(paragraph_text(output.paragraphs[1]), "乙方：公司1\t签字")
        self.assertEqual(len(output.element.xpath(".//w:br")), 1)
        self.assertEqual(len(output.element.xpath(".//w:tab")), 1)
