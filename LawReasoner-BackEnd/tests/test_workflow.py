"""Critical integration checks for auth, agent trace and contract artifacts."""

from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch
import json
import os

from docx import Document
from fastapi.testclient import TestClient

from app.api.deps import get_db
from app.config import get_settings
from app.main import app
from app.services.contract_templates import ContractTemplateCatalog


class FakeModel:
    def __init__(self, _settings):
        self.calls = 0

    async def complete(self, messages, tools):
        self.calls += 1
        if self.calls == 1:
            return {"content": None, "tool_calls": [{"id": "lookup-1", "type": "function", "function": {"name": "lookup_civil_code", "arguments": '{"article":1}'}}]}
        return {"content": "根据《民法典》第一条，需要结合具体事实分析。"}


class WorkflowTests(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        os.environ["LAW_DATABASE_PATH"] = str(Path(self.temp.name) / "app.db")
        get_settings.cache_clear()
        get_db.cache_clear()
        self.client = TestClient(app)
        self.client.__enter__()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        get_settings.cache_clear()
        get_db.cache_clear()
        os.environ.pop("LAW_DATABASE_PATH", None)
        os.environ.pop("LAW_RESOURCES_DIR", None)
        os.environ.pop("LAW_GENERATED_DIR", None)
        self.temp.cleanup()

    def test_auth_agent_trace_and_session_ownership(self):
        register = self.client.post("/api/v1/auth/register", json={"username": "alice", "email": "alice@example.org", "password": "long-password-123"})
        self.assertEqual(register.status_code, 201)
        login = self.client.post("/api/v1/auth/login", json={"username": "alice", "password": "long-password-123"})
        self.assertEqual(login.status_code, 200)
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        with patch("app.services.conversations.ChatModel", FakeModel):
            response = self.client.post("/api/v1/agent/runs", headers=headers, json={"message": "请解释民法典第一条"})
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(result["trace"][0]["result"]["article"], 1)
        self.assertEqual(result["evidence"][0]["article"], 1)
        detail = self.client.get(f"/api/v1/agent/sessions/{result['session_id']}", headers=headers)
        self.assertEqual(len(detail.json()["messages"]), 2)
        self.client.post("/api/v1/auth/register", json={"username": "bob", "email": "bob@example.org", "password": "another-password-123"})
        other = self.client.post("/api/v1/auth/login", json={"username": "bob", "password": "another-password-123"})
        other_headers = {"Authorization": f"Bearer {other.json()['access_token']}"}
        self.assertEqual(self.client.get(f"/api/v1/agent/sessions/{result['session_id']}", headers=other_headers).status_code, 404)
        self.assertEqual(self.client.post("/api/v1/auth/logout", headers=headers).status_code, 200)
        self.assertEqual(self.client.get("/api/v1/agent/sessions", headers=headers).status_code, 401)

    def test_contract_template_fields_and_generation(self):
        root = Path(self.temp.name) / "contracts"
        folder = root / "tran_template" / "Ministerial"
        placeholder_folder = folder / "占位符"
        placeholder_folder.mkdir(parents=True)
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("甲方：{")
        paragraph.add_run("甲方名称}，金额：{金额}")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "乙方：{乙方名称}"
        document.save(folder / "测试合同.docx")
        (placeholder_folder / "测试合同.json").write_text(json.dumps({"placeholders": ["甲方名称", "乙方名称", "金额"]}, ensure_ascii=False), encoding="utf-8")
        catalog = ContractTemplateCatalog(root, Path(self.temp.name) / "generated")
        template_id = catalog.list_templates()[0]["id"]
        result = catalog.generate_contract(template_id, {"甲方名称": "张三", "金额": "1000元", "无效字段": "忽略"})
        filled = Document(catalog.generated / result["filename"])
        self.assertIn("张三", filled.paragraphs[0].text)
        self.assertIn("，金额：1000元", filled.paragraphs[0].runs[1].text)
        self.assertIn("{乙方名称}", filled.tables[0].cell(0, 0).text)
        self.assertEqual(result["missing_fields"], ["乙方名称"])
        self.assertNotIn("无效字段", result["filled_fields"])

    def test_agent_requires_explicit_draft_and_download_is_user_scoped(self):
        root = Path(self.temp.name) / "contracts"
        folder = root / "tran_template" / "Ministerial"
        fields = folder / "占位符"
        fields.mkdir(parents=True)
        document = Document()
        document.add_paragraph("甲方：{甲方名称}")
        document.save(folder / "测试合同.docx")
        (fields / "测试合同.json").write_text(json.dumps({"placeholders": ["甲方名称"]}, ensure_ascii=False), encoding="utf-8")
        template_id = ContractTemplateCatalog(root, Path(self.temp.name) / "generated").list_templates()[0]["id"]
        os.environ["LAW_RESOURCES_DIR"] = str(root)
        os.environ["LAW_GENERATED_DIR"] = str(Path(self.temp.name) / "generated")
        get_settings.cache_clear()

        class DraftModel:
            def __init__(self, _settings):
                self.calls = 0

            async def complete(self, messages, tools):
                self.calls += 1
                if self.calls == 1:
                    return {"tool_calls": [{"id": "inspect", "type": "function", "function": {"name": "inspect_contract_template", "arguments": json.dumps({"template_id": template_id})}}]}
                if self.calls == 2:
                    return {"tool_calls": [{"id": "draft", "type": "function", "function": {"name": "draft_contract", "arguments": json.dumps({"template_id": template_id, "values": {"甲方名称": "张三"}}, ensure_ascii=False)}}]}
                return {"content": "合同草稿已生成，请核对。"}

        self.client.post("/api/v1/auth/register", json={"username": "alice", "email": "alice@example.org", "password": "long-password-123"})
        login = self.client.post("/api/v1/auth/login", json={"username": "alice", "password": "long-password-123"})
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        with patch("app.services.conversations.ChatModel", DraftModel):
            refused = self.client.post("/api/v1/agent/runs", headers=headers, json={"message": "请用测试合同生成合同，甲方名称张三", "allow_draft": False})
        self.assertEqual(refused.status_code, 200)
        self.assertEqual(refused.json()["artifacts"], [])
        self.assertIn("error", refused.json()["trace"][1]["result"])
        with patch("app.services.conversations.ChatModel", DraftModel):
            accepted = self.client.post("/api/v1/agent/runs", headers=headers, json={"message": "请用测试合同生成合同，甲方名称张三", "allow_draft": True})
        self.assertEqual(accepted.status_code, 200, accepted.text)
        contract_id = accepted.json()["artifacts"][0]["id"]
        self.assertEqual(self.client.get(f"/api/v1/contracts/{contract_id}/download", headers=headers).status_code, 200)
        self.assertEqual(self.client.get(f"/api/v1/contracts/{contract_id}/download").status_code, 401)


if __name__ == "__main__":
    import unittest
    unittest.main()
