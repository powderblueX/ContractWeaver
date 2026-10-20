"""Contract template catalog, retrieval and conservative DOCX filling."""

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from functools import lru_cache
import json
import re
import uuid


GROUPS = {"Ministerial": "部委", "Regional": "地方"}


@dataclass(frozen=True)
class ContractTemplate:
    id: str
    name: str
    group: str
    path: Path
    placeholder_path: Path

    def public(self) -> dict:
        return {"id": self.id, "name": self.name, "group": GROUPS[self.group]}


class ContractTemplateCatalog:
    def __init__(self, resources: Path, generated: Path, embedding_path: Path | None = None):
        self.resources = resources
        self.generated = generated
        self.embedding_path = embedding_path
        self._templates: dict[str, ContractTemplate] = {}

    def refresh(self) -> None:
        """Only files discovered beneath the configured template root can be selected."""
        templates = {}
        root = (self.resources / "tran_template").resolve()
        for group in GROUPS:
            folder = root / group
            if not folder.is_dir():
                continue
            for path in folder.glob("*.docx"):
                if not path.resolve().is_relative_to(root):
                    continue
                identifier = sha256(f"{group}/{path.stem}".encode()).hexdigest()[:20]
                templates[identifier] = ContractTemplate(
                    identifier, path.stem, group, path,
                    folder / "占位符" / f"{path.stem}.json",
                )
        self._templates = templates

    def list_templates(self, group: str | None = None, query: str = "", limit: int = 30) -> list[dict]:
        if not self._templates:
            self.refresh()
        items = self._templates.values()
        if group:
            normalized = next((key for key, label in GROUPS.items() if group in (key, label)), None)
            if normalized is None:
                raise ValueError("未知合同分类")
            items = (item for item in items if item.group == normalized)
        if query:
            items = (item for item in items if query.casefold() in item.name.casefold())
        return [item.public() for item in list(items)[:max(1, min(limit, 100))]]

    def get_template(self, template_id: str) -> ContractTemplate:
        if not self._templates:
            self.refresh()
        template = self._templates.get(template_id)
        if template is None:
            raise ValueError("合同模板不存在")
        return template

    def describe_template(self, template_id: str) -> dict:
        template = self.get_template(template_id)
        if template.placeholder_path.exists():
            payload = json.loads(template.placeholder_path.read_text(encoding="utf-8"))
            placeholders = list(dict.fromkeys(x for x in payload.get("placeholders", []) if x))
        else:
            placeholders = []
        return {**template.public(), "placeholders": placeholders}

    @staticmethod
    def _bigrams(text: str) -> set[str]:
        clean = re.sub(r"\s|[\W_]", "", text.casefold())
        return {clean[i:i + 2] for i in range(len(clean) - 1)}

    def search_templates(self, query: str, group: str | None = None, limit: int = 5) -> list[dict]:
        """Rank templates by local terms; optional FAISS improves candidate recall."""
        if not query.strip():
            raise ValueError("请描述合同需求")
        if not self._templates:
            self.refresh()
        allowed = self._templates.values()
        if group:
            normalized = next((key for key, label in GROUPS.items() if group in (key, label)), None)
            if normalized is None:
                raise ValueError("未知合同分类")
            allowed = (item for item in allowed if item.group == normalized)
        candidates = list(allowed)
        query_terms = self._bigrams(query)
        vector_candidates = self._vector_candidates(query, group)
        scored = []
        for template in candidates:
            name_terms = self._bigrams(template.name)
            keyword_terms = self._bigrams(self._keywords(template))
            name_match = len(query_terms & name_terms) / max(1, min(len(query_terms), len(name_terms)))
            keyword_match = len(query_terms & keyword_terms) / max(1, min(len(query_terms), len(keyword_terms)))
            exact = 1.0 if template.name in query or query in template.name else 0.0
            boost = vector_candidates.get((template.group, template.name), 0.0)
            score = round(0.35 * name_match + 0.2 * keyword_match + 0.15 * exact + 0.3 * boost, 3)
            if score > 0:
                scored.append((score, template))
        scored.sort(key=lambda pair: (-pair[0], pair[1].name))
        return [{**template.public(), "relevance": score} for score, template in scored[:max(1, min(limit, 20))]]

    def _keywords(self, template: ContractTemplate) -> str:
        path = self.resources / "template" / template.group / "关键词" / f"{template.name}.txt"
        return self._read_keywords(str(path))

    @staticmethod
    @lru_cache(maxsize=1024)
    def _read_keywords(path: str) -> str:
        try:
            return Path(path).read_text(encoding="utf-8")[:1000]
        except OSError:
            return ""

    def _vector_candidates(self, query: str, group: str | None) -> dict[tuple[str, str], float]:
        if not self.embedding_path:
            return {}
        try:
            import faiss
            from sentence_transformers import SentenceTransformer
            model = self._embedding_model(str(self.embedding_path))
            vector = model.encode([query]).astype("float32")
            result = {}
            for key in GROUPS:
                if group and group not in (key, GROUPS[key]):
                    continue
                folder = self.resources / "RAG" / key
                names = [Path(line.strip()).stem for line in (folder / "filenames.txt").read_text(encoding="utf-8").splitlines() if line.strip()]
                index = faiss.read_index(str(folder / "knowledge_base.index"))
                if index.ntotal != len(names):
                    continue
                distances, indices = index.search(vector, min(20, len(names)))
                for rank, position in enumerate(indices[0]):
                    if 0 <= position < len(names):
                        result[(key, names[position])] = 1.0 - rank / max(1, len(indices[0]))
            return result
        except (ImportError, OSError, ValueError, RuntimeError):
            return {}

    @staticmethod
    @lru_cache(maxsize=2)
    def _embedding_model(path: str):
        from sentence_transformers import SentenceTransformer
        return SentenceTransformer(path, local_files_only=True)

    def generate_contract(self, template_id: str, values: dict[str, str]) -> dict:
        """Fill only known placeholders; leave missing values visible for human review."""
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("服务器缺少 python-docx 依赖") from exc
        template = self.get_template(template_id)
        allowed = set(self.describe_template(template_id)["placeholders"])
        clean = {key: str(value).strip() for key, value in values.items() if key in allowed and value is not None and str(value).strip()}
        if not clean:
            raise ValueError("未提取到可填写的合同字段，请补充当事人、金额等具体信息")
        document = Document(template.path)
        filled = set()
        for paragraph in self._paragraphs(document):
            filled.update(self._fill_paragraph(paragraph, clean))
        if not filled:
            raise ValueError("模板中未找到对应占位符，无法生成合同")
        self.generated.mkdir(parents=True, exist_ok=True)
        contract_id = uuid.uuid4().hex
        filename = f"{contract_id}.docx"
        document.save(self.generated / filename)
        return {
            "id": contract_id, "template": template.public(),
            "filled_fields": sorted(filled),
            "missing_fields": sorted(allowed - filled),
            "filename": filename,
        }

    @staticmethod
    def _fill_paragraph(paragraph, values: dict[str, str]) -> set[str]:
        """Replace even split placeholders while retaining unaffected run styles."""
        runs = paragraph.runs
        original = "".join(run.text for run in runs)
        matches = [match for match in re.finditer(r"\{([^{}]+)\}", original) if match.group(1) in values]
        filled = set()

        def locate(offset: int) -> tuple[int, int]:
            cursor = 0
            for index, run in enumerate(runs):
                if offset < cursor + len(run.text):
                    return index, offset - cursor
                cursor += len(run.text)
            raise ValueError("占位符位置超出段落范围")

        # Reverse order keeps earlier character offsets stable after a replacement.
        for match in reversed(matches):
            key = match.group(1)
            first, start = locate(match.start())
            last, end = locate(match.end() - 1)
            prefix = runs[first].text[:start]
            suffix = runs[last].text[end + 1:]
            if first == last:
                runs[first].text = prefix + values[key] + suffix
            else:
                runs[first].text = prefix + values[key]
                for index in range(first + 1, last):
                    runs[index].text = ""
                runs[last].text = suffix
            filled.add(key)
        return filled

    @classmethod
    def _paragraphs(cls, document):
        for paragraph in document.paragraphs:
            yield paragraph
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cls._paragraphs(cell)
        for section in getattr(document, "sections", []):
            for part in (section.header, section.footer):
                yield from cls._paragraphs(part)
